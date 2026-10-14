"""Excel export of a user's transactions (Pro)."""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime
from typing import Iterable, List

from insights import MONTHS_NOMINATIVE

MONEY_FORMAT = '#,##0.00 "zł";-#,##0.00 "zł"'


def build_workbook(transactions: Iterable[dict], user_name: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows: List[dict] = sorted(transactions, key=lambda t: t["date"], reverse=True)

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F6B54")

    def style_header(ws, widths):
        for col, width in enumerate(widths, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22

    # Sheet 1: every transaction
    ws = wb.active
    ws.title = "Transakcje"
    ws.append(["Data", "Typ", "Kategoria", "Opis", "Kwota"])
    for t in rows:
        signed = t["amount"] if t["type"] == "income" else -t["amount"]
        ws.append([
            t["date"].date() if isinstance(t["date"], datetime) else t["date"],
            "Przychód" if t["type"] == "income" else "Wydatek",
            t.get("category") or "",
            t.get("description") or "",
            round(signed, 2),
        ])
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = "DD.MM.YYYY"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = MONEY_FORMAT
    style_header(ws, [12, 11, 22, 48, 14])
    if rows:
        ws.auto_filter.ref = f"A1:E{len(rows) + 1}"

    # Sheet 2: months
    months = defaultdict(lambda: {"income": 0.0, "expenses": 0.0})
    categories = defaultdict(lambda: defaultdict(float))
    for t in rows:
        key = (t["date"].year, t["date"].month)
        if t["type"] == "income":
            months[key]["income"] += t["amount"]
        else:
            months[key]["expenses"] += t["amount"]
            categories[key][t.get("category") or "Inne"] += t["amount"]

    ms = wb.create_sheet("Miesiące")
    ms.append(["Miesiąc", "Przychody", "Wydatki", "Bilans", "Odłożono"])
    for (year, month) in sorted(months, reverse=True):
        m = months[(year, month)]
        balance = m["income"] - m["expenses"]
        rate = balance / m["income"] if m["income"] > 0 else None
        ms.append([f"{MONTHS_NOMINATIVE[month - 1].capitalize()} {year}", round(m["income"], 2), round(m["expenses"], 2), round(balance, 2), rate])
    for row in ms.iter_rows(min_row=2, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = MONEY_FORMAT
    for row in ms.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0%"
    style_header(ms, [20, 14, 14, 14, 11])

    # Sheet 3: spending by category per month
    cs = wb.create_sheet("Kategorie")
    cs.append(["Miesiąc", "Kategoria", "Wydatki", "Udział"])
    for (year, month) in sorted(categories, reverse=True):
        total = sum(categories[(year, month)].values()) or 1
        for name, amount in sorted(categories[(year, month)].items(), key=lambda kv: -kv[1]):
            cs.append([f"{MONTHS_NOMINATIVE[month - 1].capitalize()} {year}", name, round(amount, 2), amount / total])
    for row in cs.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = MONEY_FORMAT
    for row in cs.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = "0%"
    style_header(cs, [20, 22, 14, 10])

    wb.properties.creator = "CashFlo"
    wb.properties.title = f"CashFlo – transakcje {user_name}".strip()
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
