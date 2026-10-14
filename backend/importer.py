"""Reading bank statements (CSV) and spreadsheets (Excel) into transactions.

Everything here is plain Python without the database, so it can be tested on its own.

The flow:
    table = read_table(file_bytes, file_name)        # rows of cells
    mapping, header_row, source = detect_columns(table)
    result = extract_rows(table, mapping, header_row)  # dates, amounts, descriptions
    category = suggest_category(...)                   # per row
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


class ImportFileError(Exception):
    """The file can't be read. `code` is shown to the app, which turns it into a message."""

    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

def fold(text: str) -> str:
    """Lower case without Polish letters' accents: 'Opis Operacji Ł' -> 'opis operacji l'."""
    text = (text or "").replace("ł", "l").replace("Ł", "L")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.lower()).strip()


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    return str(value).strip()


# ---------------------------------------------------------------------------
# Reading the file into a table
# ---------------------------------------------------------------------------

def read_table(content: bytes, file_name: str = "") -> List[List[str]]:
    if not content:
        raise ImportFileError("empty")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError("too_large")

    if content[:4] == b"PK\x03\x04":
        rows = _read_xlsx(content)
    elif content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        # Old .xls format: ask for .xlsx or .csv instead
        raise ImportFileError("old_excel")
    elif content[:5] in (b"%PDF-",):
        raise ImportFileError("pdf")
    else:
        rows = _read_csv(content)

    # Drop rows that are completely empty and trailing empty cells
    table: List[List[str]] = []
    for row in rows:
        cells = [cell_text(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if any(cells):
            table.append(cells)
    if not table:
        raise ImportFileError("empty")
    return table[: MAX_ROWS + 60]


def _read_xlsx(content: bytes) -> List[list]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - listed in requirements.txt
        raise ImportFileError("unreadable")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("unreadable")
    best: List[list] = []
    # Use the sheet with the most filled rows (people often keep notes on other sheets)
    for sheet in workbook.worksheets[:10]:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + 60:
                break
        filled = sum(1 for r in rows if any(v not in (None, "") for v in r))
        if filled > sum(1 for r in best if any(v not in (None, "") for v in r)):
            best = rows
    workbook.close()
    return best


def _decode(content: bytes) -> str:
    if content.startswith(b"\xff\xfe") or content.startswith(b"\xfe\xff"):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _read_csv(content: bytes) -> List[List[str]]:
    text = _decode(content)
    if "\x00" in text:
        raise ImportFileError("unreadable")
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        raise ImportFileError("empty")
    delimiter = _guess_delimiter(lines[:60])
    reader = csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter, quotechar='"')
    rows = []
    for row in reader:
        rows.append(row)
        if len(rows) > MAX_ROWS + 60:
            break
    return rows


def _guess_delimiter(lines: Sequence[str]) -> str:
    best, best_score = ",", -1.0
    for delimiter in (";", "\t", ",", "|"):
        counts = [len(r) for r in csv.reader(io.StringIO("\n".join(lines)), delimiter=delimiter, quotechar='"')]
        counts = [c for c in counts if c > 1]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        consistent = counts.count(mode)
        score = consistent * min(mode, 12)
        if score > best_score:
            best, best_score = delimiter, score
    return best


# ---------------------------------------------------------------------------
# Finding the columns
# ---------------------------------------------------------------------------

# Header words for each kind of column, most specific first. Compared after fold().
ROLE_KEYWORDS: Dict[str, List[str]] = {
    "date": [
        "data operacji", "data transakcji", "data zlecenia", "started date", "transaction date",
        "data", "date", "dzien", "kiedy", "data ksiegowania", "completed date", "data waluty",
    ],
    "amount": [
        "kwota transakcji", "kwota operacji", "kwota w walucie rachunku", "kwota", "amount",
        "wartosc", "suma", "cena", "koszt", "kwota pln", "kwota (pln)", "kwota zl", "ile", "ile zl",
    ],
    "debit": ["obciazenia", "obciazenie", "wydatki", "wydatek", "wyplaty", "debit", "money out", "paid out", "wychodzace"],
    "credit": ["uznania", "uznanie", "wplywy", "wplyw", "przychody", "przychod", "wplaty", "credit", "money in", "paid in", "przychodzace"],
    "counterparty": [
        "dane kontrahenta", "kontrahent", "nadawca / odbiorca", "nadawca/odbiorca", "odbiorca", "nadawca",
        "nazwa odbiorcy", "nazwa nadawcy", "merchant", "sklep", "gdzie", "nazwa",
    ],
    "description": [
        "opis operacji", "opis transakcji", "opis", "tytul", "tytulem", "description", "szczegoly",
        "notatka", "uwagi", "na co", "co", "komentarz", "title",
    ],
    "category": ["kategoria", "category", "rodzaj wydatku", "grupa"],
    "type": ["typ", "typ transakcji", "rodzaj transakcji", "type", "rodzaj", "przychod/wydatek", "wydatek/przychod"],
    "currency": ["waluta", "currency"],
    "status": ["status", "state"],
    "balance": ["saldo", "saldo po transakcji", "saldo po operacji", "balance", "saldo ksiegowe"],
}

ROLES_SINGLE = ("date", "amount", "debit", "credit", "category", "type", "currency", "status")


@dataclass
class Mapping:
    date: Optional[int] = None
    amount: Optional[int] = None
    debit: Optional[int] = None
    credit: Optional[int] = None
    description: List[int] = field(default_factory=list)
    category: Optional[int] = None
    type: Optional[int] = None
    currency: Optional[int] = None
    status: Optional[int] = None

    def as_dict(self) -> dict:
        return {
            "date": self.date,
            "amount": self.amount,
            "debit": self.debit,
            "credit": self.credit,
            "description": list(self.description),
            "category": self.category,
            "type": self.type,
        }

    @classmethod
    def from_dict(cls, data: dict, columns: int) -> "Mapping":
        def idx(key):
            value = data.get(key)
            if isinstance(value, int) and 0 <= value < columns:
                return value
            return None

        desc = [d for d in (data.get("description") or []) if isinstance(d, int) and 0 <= d < columns]
        return cls(
            date=idx("date"),
            amount=idx("amount"),
            debit=idx("debit"),
            credit=idx("credit"),
            description=desc[:6],
            category=idx("category"),
            type=idx("type"),
            currency=idx("currency"),
            status=idx("status"),
        )

    def usable(self) -> bool:
        return self.date is not None and (self.amount is not None or self.debit is not None or self.credit is not None)


# Columns that hold account numbers, bank names etc. are never the description,
# and blocked or converted amounts are never the amount.
_NOT_TEXT = ("bank", "rachun", "konto", "konta", "numer", "nr ", "iban", "referenc", "saldo", "balance")
_NOT_AMOUNT = ("blokad", "saldo", "balance", "kurs ", "numer", "nr ")
_TEXT_ROLES = ("counterparty", "description", "category", "type")


def _header_role(cell: str) -> Optional[Tuple[str, int]]:
    """Which kind of column a header names, and how sure we are (lower = better)."""
    name = fold(cell).lstrip("#").strip(" :*")
    if not name or len(name) > 60:
        return None
    padded = name + " "
    not_text = any(x in padded for x in _NOT_TEXT)
    not_amount = any(x in padded for x in _NOT_AMOUNT)
    for phase in (0, 1):
        for role, words in ROLE_KEYWORDS.items():
            if not_text and role in _TEXT_ROLES:
                continue
            if not_amount and role in ("amount", "debit", "credit"):
                continue
            for rank, word in enumerate(words):
                if phase == 0 and name == word:
                    return role, rank
                if phase == 1 and len(word) >= 4 and any(name.startswith(word + sep) for sep in (" ", "(", "/")):
                    return role, rank + 20
    return None


def find_header(table: List[List[str]]) -> Tuple[Optional[int], Mapping]:
    """Find the header row in the first rows of the table and map its columns."""
    best_row, best_mapping, best_score = None, Mapping(), 0
    for row_index, row in enumerate(table[:40]):
        roles: Dict[str, List[Tuple[int, int]]] = {}
        for col, cell in enumerate(row):
            found = _header_role(cell)
            if found:
                role, rank = found
                roles.setdefault(role, []).append((rank, col))
        mapping = Mapping()
        for role in ROLES_SINGLE:
            if role in roles:
                setattr(mapping, role, sorted(roles[role])[0][1])
        # Merchant name first, then the title / description
        description_cols = [c for _, c in sorted(roles.get("counterparty", []))] + [
            c for _, c in sorted(roles.get("description", []))
        ]
        mapping.description = list(dict.fromkeys(description_cols))[:4]
        score = sum(1 for r in ("date", "amount", "debit", "credit", "category", "type") if getattr(mapping, r) is not None)
        score += 1 if mapping.description else 0
        if mapping.usable() and score > best_score:
            best_row, best_mapping, best_score = row_index, mapping, score
    if best_row is None:
        return None, Mapping()

    # Bank exports (e.g. PKO) put details in columns without a header after the description
    header = table[best_row]
    width = max((len(r) for r in table[best_row + 1: best_row + 30]), default=len(header))
    if best_mapping.description:
        last = max(best_mapping.description)
        extra = [c for c in range(last + 1, width) if c >= len(header) or not header[c].strip()]
        best_mapping.description = (best_mapping.description + extra)[:6]
    return best_row, best_mapping


def guess_columns_by_content(table: List[List[str]]) -> Mapping:
    """For files without a header row: find the date, amount and text columns by what they contain."""
    rows = table[:200]
    width = max((len(r) for r in rows), default=0)
    stats = []
    for col in range(width):
        values = [r[col] for r in rows if col < len(r) and r[col].strip()]
        if not values:
            stats.append((0.0, 0.0, 0.0, False))
            continue
        dates = sum(1 for v in values if parse_date(v) is not None) / len(values)
        amounts = sum(1 for v in values if parse_amount(v) is not None and parse_date(v) is None) / len(values)
        negatives = any((parse_amount(v) or 0) < 0 for v in values)
        text = sum(len(v) for v in values if parse_amount(v) is None and parse_date(v) is None) / len(values)
        stats.append((dates, amounts, text, negatives))

    mapping = Mapping()
    date_cols = [c for c, s in enumerate(stats) if s[0] >= 0.6]
    if date_cols:
        mapping.date = date_cols[0]
    amount_cols = [c for c, s in enumerate(stats) if s[1] >= 0.6 and c != mapping.date]
    if amount_cols:
        # Prefer a column with negative numbers (spending), else the first numeric one
        with_negatives = [c for c in amount_cols if stats[c][3]]
        mapping.amount = (with_negatives or amount_cols)[0]
    text_cols = sorted(
        (c for c, s in enumerate(stats) if s[2] >= 3 and c not in (mapping.date, mapping.amount)),
        key=lambda c: -stats[c][2],
    )
    mapping.description = text_cols[:2]
    return mapping


def detect_source(table: List[List[str]], header_row: Optional[int], is_excel: bool) -> str:
    head = " | ".join(fold(c) for r in table[: (header_row or 0) + 1] for c in r)
    if "#data operacji" in head or "mbank" in head:
        return "mBank"
    if "dane kontrahenta" in head:
        return "ING"
    if "saldo po transakcji" in head or ("opis transakcji" in head and "typ transakcji" in head):
        return "PKO BP"
    if "started date" in head and "completed date" in head:
        return "Revolut"
    if "nadawca / odbiorca" in head or "kwota operacji" in head:
        return "Pekao"
    if "obciazenia" in head and "uznania" in head:
        return "Millennium"
    if "santander" in head:
        return "Santander"
    return "Excel" if is_excel else "CSV"


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------

_CURRENCY_RE = re.compile(r"(pln|zl|zł|eur|usd|gbp|chf|€|\$|£)", re.IGNORECASE)


def parse_amount(value) -> Optional[float]:
    """'-1 234,56 PLN', '1.234,56', '(12.50)', '12,50-' -> signed float. None if not a number."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    if not text or len(text) > 30:
        return None
    text = _CURRENCY_RE.sub("", text)
    text = text.replace("−", "-").replace("–", "-").replace("\xa0", "").replace(" ", "")
    text = text.replace(" ", "").replace("'", "")
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    if text.endswith("-"):
        negative, text = True, text[:-1]
    if text.startswith("+"):
        text = text[1:]
    if text.startswith("-"):
        negative, text = not negative, text[1:]
    if not text or not re.fullmatch(r"[\d.,]+", text) or not re.search(r"\d", text):
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")  # 1.234,56
        else:
            text = text.replace(",", "")  # 1,234.56
    elif "," in text:
        parts = text.split(",")
        if len(parts) > 2 and all(len(p) == 3 for p in parts[1:]):
            text = text.replace(",", "")  # 1,234,567
        elif len(parts) == 2:
            text = text.replace(",", ".")  # 12,50 (Polish decimal comma)
        else:
            return None
    elif text.count(".") > 1:
        parts = text.split(".")
        if all(len(p) == 3 for p in parts[1:]):
            text = text.replace(".", "")  # 1.234.567
        else:
            return None
    try:
        number = float(text)
    except ValueError:
        return None
    return -number if negative else number


_DATE_FORMATS = (
    "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%Y.%m.%d", "%d/%m/%Y", "%Y/%m/%d",
    "%d.%m.%y", "%d-%m-%y", "%Y%m%d",
)


def parse_date(value) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text or len(text) > 40:
        return None
    first = re.split(r"[ T]", text)[0]
    for fmt in _DATE_FORMATS:
        try:
            parsed = datetime.strptime(first, fmt).date()
        except ValueError:
            continue
        if fmt == "%Y%m%d" and not (1990 <= parsed.year <= 2100):
            return None
        if 1990 <= parsed.year <= 2100:
            return parsed
    # Excel serial number (days since 1899-12-30)
    if re.fullmatch(r"\d{5}(\.\d+)?", text):
        serial = float(text)
        if 32874 <= serial <= 73051:  # 1990 .. 2100
            return date(1899, 12, 30) + timedelta(days=int(serial))
    return None


_NOISE_PATTERNS = [
    r"\b\d{4}[ -]?\*{2,}[ -]?\*{0,4}[ -]?\d{4}\b",  # card numbers 4251 **** **** 1234
    r"\b\d{6}\*+\d{4}\b",
    r"\b(nr|numer)\s+(karty|rachunku|ref\.?|referencyjny|transakcji)[:\s]*[\w*/-]+",
    r"\boryginalna kwota operacji[:\s]*[\d.,]+\s*\w*",
    r"\bdata (wykonania|transakcji|operacji|ksiegowania|waluty)[:\s]*[\d./-]+",
    r"\b\d{4}-\d{2}-\d{2}\b",
    r"\b\d{2}[./]\d{2}[./]\d{2,4}\b",
    r"\b\d{2}:\d{2}(:\d{2})?\b",
    r"\b\d{3,}\b",
    r"\b[a-z]{2}\d{2}(\s?\d{4}){4,7}\b",  # IBAN
]
_LABELS = [
    "numer karty", "nr karty", "data wykonania", "oryginalna kwota operacji", "numer rachunku", "nr rachunku",
    "nazwa odbiorcy", "nazwa nadawcy", "adres odbiorcy", "adres nadawcy", "rachunek odbiorcy", "rachunek nadawcy",
    "tytul", "lokalizacja", "adres", "miasto", "kraj", "numer telefonu", "operacja", "numer referencyjny",
    "dane kontrahenta", "odbiorca", "nadawca", "opis",
]
_ACCENTS = {"a": "[aą]", "c": "[cć]", "e": "[eę]", "l": "[lł]", "n": "[nń]", "o": "[oó]", "s": "[sś]", "z": "[zźż]"}


def _accent_pattern(word: str) -> str:
    return "".join(_ACCENTS.get(ch, re.escape(ch)) for ch in word)


_LABEL_RE = re.compile(r"(?i)\b(" + "|".join(_accent_pattern(l) for l in _LABELS) + r")\s*:")
_KEEP_LABELS = ("adres", "nazwa odbiorcy", "nazwa nadawcy", "tytul", "odbiorca", "nadawca", "dane kontrahenta", "opis")


def clean_description(parts: Iterable[str]) -> str:
    """Make a bank's description short and readable: 'ZAKUP KARTĄ BIEDRONKA 1234 WARSZAWA' -> 'Biedronka Warszawa'."""
    pieces: List[str] = []
    for part in parts:
        part = (part or "").strip().strip('"').strip()
        if not part:
            continue
        # "Lokalizacja: Adres: NETFLIX.COM Miasto: Amsterdam Kraj: Holandia" -> keep useful labelled values
        labelled = _LABEL_RE.split(part)
        if len(labelled) > 1:
            useful = []
            for i in range(1, len(labelled) - 1, 2):
                label = fold(labelled[i])
                value = labelled[i + 1].strip(" ,;")
                if label in _KEEP_LABELS and value:
                    useful.append(value)
            lead = labelled[0].strip(" ,;")
            if lead:
                useful.insert(0, lead)
            part = " ".join(useful)
        pieces.append(part)

    text = " ".join(pieces)
    for pattern in _NOISE_PATTERNS:
        text = re.sub(pattern, " ", text, flags=re.IGNORECASE)
    for phrase in (
        "zakup przy uzyciu karty", "zakup przy użyciu karty", "platnosc karta", "płatność kartą", "transakcja karta",
        "transakcja kartą", "zakup bilety", "platnosc blik", "płatność blik", "przelew blik", "przelew na telefon",
        "przelew wychodzacy", "przelew wychodzący", "przelew przychodzacy", "przelew przychodzący", "przelew zewnetrzny",
        "przelew zewnętrzny", "przelew wewnetrzny", "przelew wewnętrzny", "transakcja bezgotowkowa",
        "transakcja bezgotówkowa", "card payment", "zakup kartą", "zakup karta", "płatność web blik", "platnosc web blik",
    ):
        text = re.sub(re.escape(phrase), " ", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*[|;]\s*", " ", text)
    text = re.sub(r"\s+", " ", text).strip(" ,.-/:*#")
    if len(text) > 4 and text.upper() == text:
        text = smart_title(text)
    return text[:120]


_SMALL_WORDS = {"za", "na", "do", "od", "i", "w", "z", "ze", "po", "we", "dla", "u", "o", "a"}


def smart_title(text: str) -> str:
    """'BIEDRONKA WARSZAWA' -> 'Biedronka Warszawa', keeping short codes like 'BP' or 'PKP'."""
    words = []
    for word in text.split(" "):
        if word.lower() in _SMALL_WORDS:
            words.append(word.lower())
        elif len(word) <= 3 and word.isalpha():
            words.append(word)
        else:
            words.append(word[:1].upper() + word[1:].lower())
    return " ".join(words)


_INCOME_WORDS = ("przychod", "przychód", "wplyw", "wpływ", "uznanie", "income", "credit", "dochod", "dochód", "wpłata", "wplata")
_EXPENSE_WORDS = ("wydatek", "wydatki", "obciazenie", "obciążenie", "expense", "debit", "koszt", "platnosc", "płatność")
_SKIP_STATUSES = ("reverted", "declined", "failed", "odrzucon", "anulowan", "cancelled", "canceled", "pending", "oczekuj")


@dataclass
class ExtractedRow:
    index: int
    date: date
    amount: float  # always positive
    type: str  # "expense" | "income"
    description: str
    category_hint: str = ""


@dataclass
class ExtractResult:
    rows: List[ExtractedRow]
    skipped: int = 0
    foreign_currency: int = 0
    assumed_expenses: bool = False


def extract_rows(table: List[List[str]], mapping: Mapping, header_row: Optional[int]) -> ExtractResult:
    start = (header_row + 1) if header_row is not None else 0
    rows: List[ExtractedRow] = []
    skipped = 0
    foreign = 0

    def cell(row, col):
        if col is None or col >= len(row):
            return ""
        return row[col]

    raw = []
    for index in range(start, len(table)):
        row = table[index]
        when = parse_date(cell(row, mapping.date))
        if when is None:
            skipped += 1 if any(c.strip() for c in row) else 0
            continue

        status = fold(cell(row, mapping.status))
        if status and any(s in status for s in _SKIP_STATUSES):
            skipped += 1
            continue
        currency = fold(cell(row, mapping.currency))
        if currency and currency not in ("pln", "zl", "zł"):
            foreign += 1
            continue

        signed: Optional[float] = None
        explicit_type: Optional[str] = None
        if mapping.amount is not None:
            signed = parse_amount(cell(row, mapping.amount))
        if signed in (None, 0) and (mapping.debit is not None or mapping.credit is not None):
            debit = parse_amount(cell(row, mapping.debit)) if mapping.debit is not None else None
            credit = parse_amount(cell(row, mapping.credit)) if mapping.credit is not None else None
            if debit:
                signed, explicit_type = -abs(debit), "expense"
            elif credit:
                signed, explicit_type = abs(credit), "income"
        if signed is None or signed == 0 or abs(signed) >= 100_000_000:
            skipped += 1
            continue

        type_text = fold(cell(row, mapping.type))
        if explicit_type is None and type_text:
            if any(w in type_text for w in (fold(x) for x in _INCOME_WORDS)):
                explicit_type = "income"
            elif any(w in type_text for w in (fold(x) for x in _EXPENSE_WORDS)):
                explicit_type = "expense"

        description = clean_description(cell(row, c) for c in mapping.description)
        hint = cell(row, mapping.category).strip()[:60]
        raw.append((index, when, signed, explicit_type, description, hint))

    # A spreadsheet with only positive amounts and no other clue: treat them as expenses
    has_negative = any(r[2] < 0 for r in raw)
    has_explicit = any(r[3] for r in raw)
    assumed = bool(raw) and not has_negative and not has_explicit and mapping.debit is None and mapping.credit is None

    for index, when, signed, explicit_type, description, hint in raw:
        if explicit_type:
            tx_type = explicit_type
        elif assumed:
            tx_type = "expense"
        else:
            tx_type = "income" if signed > 0 else "expense"
        rows.append(ExtractedRow(index, when, round(abs(signed), 2), tx_type, description, hint))
        if len(rows) >= MAX_ROWS:
            break

    return ExtractResult(rows=rows, skipped=skipped, foreign_currency=foreign, assumed_expenses=assumed)


# ---------------------------------------------------------------------------
# Categories
# ---------------------------------------------------------------------------

# Keywords in a description (or the bank's own category) -> default CashFlo category
CATEGORY_RULES: List[Tuple[str, List[str]]] = [
    ("Jedzenie i Picie", [
        "biedronka", "lidl", "zabka", "kaufland", "auchan", "carrefour", "netto", "dino", "aldi", "stokrotka",
        "lewiatan", "delikatesy", "spar", "polomarket", "intermarche", "mcdonald", "kfc", "burger", "pizza",
        "restaurac", "bistro", "kebab", "sushi", "pyszne", "glovo", "wolt", "uber eats", "ubereats", "starbucks",
        "costa coffee", "kawiarni", "piekarni", "cukierni", "freshmarket", "groszek", "chata polska", "mila ",
        "makro", "selgros", "frisco", "piotr i pawel", "zywnosc", "spozywcz", "jedzenie", "bar ", "food",
        "market", "supermarket", "grocer", "lody", "kawa", "coffee",
    ]),
    ("Transport", [
        "orlen", "bp ", "shell", "circle k", "lotos", "moya", "amic", "jakdojade", "mpk", "ztm", "pkp", "intercity",
        "koleje", "polregio", "bolt", "freenow", "free now", "taxi", "parking", "autostrad", "viatoll", "e-toll",
        "flixbus", "ryanair", "wizz", "lot polish", "paliwo", "stacja", "myjnia", "warsztat", "opony", "uber",
        "transport", "skm", "metro", "hulajnog", "traficar", "panek", "car wash",
    ]),
    ("Rozrywka", [
        "netflix", "spotify", "hbo", "max.com", "disney", "youtube", "canal+", "canal plus", "player.pl", "polsat box",
        "apple music", "tidal", "deezer", "steam", "playstation", "xbox", "nintendo", "kino", "cinema city",
        "multikino", "helios", "teatr", "koncert", "ebilet", "going.", "audioteka", "storytel", "legimi", "skyshowtime",
        "twitch", "rozrywk", "gry ", "games", "muzeum", "zoo", "kregiel", "bowling",
    ]),
    ("Zdrowie", [
        "apteka", "doz", "gemini", "super-pharm", "superpharm", "medicover", "luxmed", "lux med", "enel-med",
        "enelmed", "przychodni", "lekarz", "stomatolog", "dentyst", "szpital", "laborator", "diagnostyka", "alab",
        "optyk", "silownia", "fitness", "multisport", "medycz", "zdrowi", "rehabilit", "fizjoterap", "pharma",
    ]),
    ("Edukacja", [
        "szkola", "przedszkol", "kurs", "udemy", "coursera", "uczelni", "uniwersytet", "czesne", "korepetyc",
        "duolingo", "edukacj", "ksiegarni", "podrecznik", "babbel",
    ]),
    ("Zakupy", [
        "allegro", "amazon", "zalando", "aliexpress", "temu", "shein", "rossmann", "hebe", "douglas", "sephora",
        "ccc", "deichmann", "reserved", "sinsay", "house", "cropp", "mohito", "h&m", "zara", "pepco", "action",
        "tk maxx", "half price", "media expert", "euro rtv", "rtv euro", "x-kom", "morele", "empik", "smyk",
        "decathlon", "4f", "ebay", "olx", "vinted", "zakupy", "odziez", "ubrania", "kosmetyk", "elektronik",
        "drogeri", "apple.com", "google play", "shop", "sklep",
    ]),
    ("Dom i Rachunki", [
        "czynsz", "pge", "tauron", "enea", "energa", "innogy", "e.on", "eon ", "pgnig", "fortum", "veolia",
        "wodociag", "mpwik", "gaz ", "prad", "energia", "orange", "play ", "p4 ", "plus ", "polkomtel", "t-mobile",
        "netia", "upc", "vectra", "inea", "internet", "telefon", "abonament", "leroy merlin", "castorama", "obi ",
        "ikea", "jysk", "bricomarche", "mrowka", "wspolnota", "spoldzielni", "administracj", "mieszkan", "rachun",
        "media", "smieci", "ubezpiecz", "pzu", "warta", "allianz", "ergo hestia", "link4", "dom ", "home",
    ]),
]

INCOME_RULES: List[Tuple[str, List[str]]] = [
    ("Wypłata", ["wynagrodzenie", "pensja", "pensji", "wyplata", "wypłata", "salary", "payroll", "pobory", "uposazenie"]),
    ("Premia", ["premia", "bonus", "nagroda", "trzynastka"]),
]

SUBSCRIPTION_BRANDS = {
    "netflix": "Netflix", "spotify": "Spotify", "hbo": "HBO Max", "max.com": "Max", "disney": "Disney+",
    "youtube": "YouTube Premium", "canal+": "Canal+", "player.pl": "Player", "polsat box": "Polsat Box",
    "apple.com": "Apple", "icloud": "iCloud", "apple music": "Apple Music", "google": "Google",
    "tidal": "Tidal", "deezer": "Deezer", "audioteka": "Audioteka", "storytel": "Storytel", "legimi": "Legimi",
    "empik go": "Empik Go", "skyshowtime": "SkyShowtime", "amazon prime": "Amazon Prime", "prime video": "Prime Video",
    "xbox": "Xbox", "playstation": "PlayStation", "nintendo": "Nintendo", "chatgpt": "ChatGPT", "openai": "ChatGPT",
    "adobe": "Adobe", "microsoft": "Microsoft 365", "dropbox": "Dropbox", "multisport": "Multisport",
    "medicover": "Medicover", "luxmed": "LuxMed", "duolingo": "Duolingo", "babbel": "Babbel",
    "tinder": "Tinder", "bumble": "Bumble", "patreon": "Patreon", "twitch": "Twitch", "allegro smart": "Allegro Smart",
}


def suggest_category(
    description: str,
    tx_type: str,
    hint: str,
    categories: Sequence[dict],
) -> str:
    """Pick one of the user's categories for an imported row."""
    of_type = [c for c in categories if (c.get("type") or "expense") == tx_type]
    by_name = {fold(c["name"]): c["name"] for c in of_type}

    # 1. The file names a category the user already has
    if hint and fold(hint) in by_name:
        return by_name[fold(hint)]

    text = " " + fold(f"{description} {hint}") + " "
    # 2. One of the user's own categories is named in the text ("Paliwo", "Kot")
    for c in of_type:
        folded = fold(c["name"])
        if not c.get("is_default") and len(folded) >= 3 and f" {folded}" in text:
            return c["name"]
    # 3. Built-in rules for well-known shops and services
    rules = INCOME_RULES if tx_type == "income" else CATEGORY_RULES
    for name, words in rules:
        if any(fold(word) in text for word in words):
            if fold(name) in by_name:
                return by_name[fold(name)]

    fallback = "Inne przychody" if tx_type == "income" else "Inne"
    if fold(fallback) in by_name:
        return by_name[fold(fallback)]
    return of_type[0]["name"] if of_type else fallback


def refine_assumed_types(result: "ExtractResult", categories: Sequence[dict]) -> None:
    """When every amount was positive (a home budget sheet), mark rows that look like income."""
    if not result.assumed_expenses:
        return
    income_names = {fold(c["name"]) for c in categories if (c.get("type") or "expense") == "income"}
    expense_names = {fold(c["name"]) for c in categories if (c.get("type") or "expense") == "expense"}
    income_words = [fold(w) for _, words in INCOME_RULES for w in words] + ["przychod", "zwrot podatku"]
    for row in result.rows:
        hint = fold(row.category_hint)
        text = " " + fold(f"{row.description} {row.category_hint}") + " "
        if (hint and hint in income_names and hint not in expense_names) or any(w in text for w in income_words):
            row.type = "income"


# ---------------------------------------------------------------------------
# Duplicates
# ---------------------------------------------------------------------------

def description_key(description: str) -> str:
    return re.sub(r"[^a-z0-9]", "", fold(description))[:40]


def row_hash(when: date, amount: float, tx_type: str, description: str, occurrence: int) -> str:
    base = f"{when.isoformat()}|{amount:.2f}|{tx_type}|{description_key(description)}|{occurrence}"
    return hashlib.sha1(base.encode()).hexdigest()


def hashes_for(rows: Sequence[ExtractedRow]) -> List[str]:
    """One hash per row; identical rows in one file (two coffees) get different hashes."""
    seen: Dict[str, int] = {}
    out = []
    for r in rows:
        base = f"{r.date.isoformat()}|{r.amount:.2f}|{r.type}|{description_key(r.description)}"
        n = seen.get(base, 0)
        seen[base] = n + 1
        out.append(row_hash(r.date, r.amount, r.type, r.description, n))
    return out


# ---------------------------------------------------------------------------
# Whole file
# ---------------------------------------------------------------------------

@dataclass
class ParsedFile:
    table: List[List[str]]
    header_row: Optional[int]
    mapping: Mapping
    source: str
    result: ExtractResult

    @property
    def columns(self) -> List[str]:
        width = max((len(r) for r in self.table[: (self.header_row or 0) + 30]), default=0)
        if self.header_row is not None:
            header = self.table[self.header_row]
            return [(header[c] if c < len(header) and header[c] else f"Kolumna {c + 1}") for c in range(width)]
        return [f"Kolumna {c + 1}" for c in range(width)]

    def sample(self, count: int = 5) -> List[List[str]]:
        start = (self.header_row + 1) if self.header_row is not None else 0
        return [row[:12] for row in self.table[start: start + count]]


def parse_file(content: bytes, file_name: str = "", mapping: Optional[Mapping] = None,
               table: Optional[List[List[str]]] = None, header_row: Optional[int] = None) -> ParsedFile:
    is_excel = content[:4] == b"PK\x03\x04" if content else file_name.lower().endswith(".xlsx")
    if table is None:
        table = read_table(content, file_name)
    if mapping is None:
        header_row, mapping = find_header(table)
        if header_row is None:
            mapping = guess_columns_by_content(table)
    source = detect_source(table, header_row, is_excel)
    result = extract_rows(table, mapping, header_row) if mapping.usable() else ExtractResult(rows=[])
    return ParsedFile(table=table, header_row=header_row, mapping=mapping, source=source, result=result)
