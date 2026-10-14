"""Excel export opens and holds the right numbers."""

import io
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from exporter import build_workbook  # noqa: E402


def test_workbook_contents():
    from openpyxl import load_workbook

    data = build_workbook([
        {"date": datetime(2026, 9, 1, 12), "type": "income", "amount": 5000, "category": "Wypłata", "description": "Pensja"},
        {"date": datetime(2026, 9, 3, 12), "type": "expense", "amount": 1800, "category": "Dom i Rachunki", "description": "Czynsz"},
        {"date": datetime(2026, 10, 2, 12), "type": "expense", "amount": 45.99, "category": "Jedzenie i Picie", "description": "Biedronka"},
    ], "Dawid")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Transakcje", "Miesiące", "Kategorie"]
    tx = list(wb["Transakcje"].values)
    assert tx[0] == ("Data", "Typ", "Kategoria", "Opis", "Kwota")
    assert tx[1][1:] == ("Wydatek", "Jedzenie i Picie", "Biedronka", -45.99)
    months = list(wb["Miesiące"].values)
    assert months[1][0] == "Październik 2026"
    assert months[2][:4] == ("Wrzesień 2026", 5000, 1800, 3200)
    assert round(months[2][4], 2) == 0.64


if __name__ == "__main__":
    test_workbook_contents()
    print("ok   test_workbook_contents")
